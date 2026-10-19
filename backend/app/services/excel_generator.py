
import unicodedata
import warnings
from pathlib import Path
from typing import Any, Dict
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl import Workbook
import pandas as pd
from pandas.api.types import is_datetime64_any_dtype

from app.services.excel_reader import read_excel_path


def _parse_dates(obj):
    """pd.to_datetime en dayfirst, valeurs invalides → NaT, sans le UserWarning
    « Could not infer format » (les dates textuelles ont des formats variés)."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return pd.to_datetime(obj, dayfirst=True, errors="coerce")


# Conversion d'un format strftime Python (ex. "%d/%m/%Y %H:%M:%S") en format de
# nombre Excel (ex. "DD/MM/YYYY HH:MM:SS") pour écrire de VRAIES dates typées,
# affichées comme l'utilisateur le souhaite (et non du texte).
_STRFTIME_TO_EXCEL = [
    ("%Y", "YYYY"), ("%y", "YY"),
    ("%m", "MM"), ("%d", "DD"),
    ("%H", "HH"), ("%I", "HH"),
    ("%M", "MM"), ("%S", "SS"),
    ("%p", "AM/PM"),
]


def _excel_number_format(strftime_fmt: str) -> str:
    """Traduit un format strftime en format de nombre Excel.
    Excel distingue mois/minutes par le contexte (« MM » après « HH: » = minutes)."""
    fmt = strftime_fmt or "%d/%m/%Y"
    for token, repl in _STRFTIME_TO_EXCEL:
        fmt = fmt.replace(token, repl)
    return fmt


def _is_date_column(series: "pd.Series") -> bool:
    """Détecte une colonne de dates, y compris quand les dates sont stockées en
    TEXTE (exports Jira/HTML). Conservateur : exige un séparateur de date et un
    fort taux de valeurs réellement analysables comme dates (évite les faux
    positifs sur des identifiants ou des nombres)."""
    if is_datetime64_any_dtype(series):
        return True
    if series.dtype != object:
        return False
    sample = series.dropna()
    if sample.empty:
        return False
    sample = sample.head(200).astype(str).str.strip()
    # Ressemble à une date : contient un chiffre ET un séparateur / ou -
    date_like = sample.str.contains(r"\d") & sample.str.contains(r"[/\-]")
    if date_like.mean() < 0.8:
        return False
    parsed = _parse_dates(sample)
    return parsed.notna().mean() >= 0.8


def _normalize(value: Any) -> str:
    """Minuscule, sans accents ni espaces superflus — pour comparer libellés/valeurs."""
    s = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return s.lower().strip()


def _first_word(value: Any) -> Any:
    """Premier mot d'un champ texte.
    Ex. « Production – Exploitation » → « Production », « Validation – Validation » → « Validation ».
    Les valeurs vides/NaN sont conservées telles quelles."""
    if value is None:
        return value
    try:
        if pd.isna(value):
            return value
    except (TypeError, ValueError):
        pass
    parts = str(value).strip().split()
    return parts[0] if parts else value


# Colonnes considérées comme "Priorité" et valeurs considérées comme "Critique"
PRIORITY_COLUMN_NAMES = {"priorite", "priority"}
CRITICAL_VALUES = {"critique", "critical"}

# Colonne dont on ne garde que le premier mot de chaque valeur
ENV_PHASE_COLUMN_NAMES = {"environnement et phase"}

# Colonne source contenant la date de création du ticket (utilisée pour la colonne "Year")
CREATION_COLUMN_NAMES = {"creation"}

# Colonnes dont les valeurs gardent leur alignement à gauche (toutes les autres sont centrées)
LEFT_ALIGN_COLUMN_NAMES = {"cle", "resume"}

class ExcelGenerator:
    def __init__(self):
        # Styles pour les en-têtes et les cellules
        self.header_front = Font(bold=True, size=10, name="Arial")
        self.header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
        self.header_font_color = Font(bold=True, size=10, name="Arial", color="000000")

        self.cell_font = Font(size=9, name="Arial")
        # Police rouge en gras pour les valeurs "Critique"
        self.critical_font = Font(size=9, name="Arial", color="CC0000", bold=True)
        self.center_align = Alignment(horizontal="center", vertical="center")
        # Centré avec retour à la ligne (en-têtes) — limite la largeur des cellules
        self.center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

        self.border = Border(
            left=Side(border_style="thin", color="000000"),
            right=Side(border_style="thin", color="000000"),
            top=Side(border_style="thin", color="000000"),
            bottom=Side(border_style="thin", color="000000")
        )

    async def generate_excel(
        self,
        input_path: Path,
        output_path: Path,
        columns: list[dict[str, Any]],
        date_format: str = "%d/%m/%Y %H:%M:%S"
    ) -> Dict[str, Any]:
        """Génère un fichier Excel à partir des données analysées"""

        # Lire les données d'entrée (en sautant les lignes de préambule éventuelles)
        df = read_excel_path(input_path)
        df.columns = [str(c).strip() for c in df.columns]

        # Extraire les noms des colonnes sélectionnées dans l'ordre demandé
        selected_columns = [col["name"] for col in columns]

        # Filtrer le DataFrame pour ne garder que les colonnes sélectionnées
        # et dans l'ordre spécifié
        df_filtered = df[selected_columns]

        # Calculer les largeurs sur les valeurs RÉELLEMENT affichées (après transformations)
        # afin que les colonnes "Environnement et phase" (premier mot) et "Création"
        # (date/heure compacte) ne soient pas surdimensionnées.
        df_for_width = df_filtered.copy()
        for col in selected_columns:
            if _normalize(col) in ENV_PHASE_COLUMN_NAMES:
                df_for_width[col] = df_for_width[col].map(_first_word)
            elif _normalize(col) in CREATION_COLUMN_NAMES:
                df_for_width[col] = "00/00/0000"  # une ligne (date), largeur compacte
        column_widths = self.calculate_column_widths(df_for_width, selected_columns)

        # Créer un nouveau classeur Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Analyse des données"

        # Repérer l'index (1-based) de la colonne "Priorité" si elle est sélectionnée
        priority_col_idx = None
        for col_idx, col_info in enumerate(columns, start=1):
            if _normalize(col_info["name"]) in PRIORITY_COLUMN_NAMES:
                priority_col_idx = col_idx
                break

        # Repérer l'index (1-based) de la colonne "Environnement et phase" si elle est sélectionnée
        env_phase_col_idx = None
        for col_idx, col_info in enumerate(columns, start=1):
            if _normalize(col_info["name"]) in ENV_PHASE_COLUMN_NAMES:
                env_phase_col_idx = col_idx
                break

        # Colonne "Year" ajoutée en fin de tableau : année de création du ticket
        # (déduite de la colonne "Création" de la source, même si elle n'est pas sélectionnée)
        creation_col = next(
            (c for c in df.columns if _normalize(c) in CREATION_COLUMN_NAMES), None
        )
        years = None
        year_col_idx = None
        if creation_col is not None:
            parsed = _parse_dates(df[creation_col])
            years = [None if pd.isna(d) else int(d.year) for d in parsed]
            year_col_idx = len(columns) + 1

        # Index (1-based) de la colonne "Création" parmi les colonnes sélectionnées
        # (sa date est affichée sur deux lignes : date puis heure)
        creation_col_idx = None
        for col_idx, col_info in enumerate(columns, start=1):
            if _normalize(col_info["name"]) in CREATION_COLUMN_NAMES:
                creation_col_idx = col_idx
                break

        # Index (1-based) des colonnes à laisser alignées à gauche ("Clé", "Résumé")
        left_align_col_idxs = {
            col_idx
            for col_idx, col_info in enumerate(columns, start=1)
            if _normalize(col_info["name"]) in LEFT_ALIGN_COLUMN_NAMES
        }

        # Format de date Excel dérivé du format demandé, pour écrire de VRAIES dates
        # (typées, filtrables/triables) plutôt que du texte.
        excel_date_format = _excel_number_format(date_format)
        # Colonne "Création" : date et heure sur deux lignes (saut de ligne dans le
        # format de nombre + wrap_text) → reste compacte tout en restant une date.
        excel_date_format_creation = excel_date_format.replace(" ", "\n")

        # Index (1-based) des colonnes de type date (hors "Création", traitée à part) :
        # leurs valeurs seront écrites comme dates réelles avec le format d'affichage.
        date_col_idxs = set()
        for col_idx, col_info in enumerate(columns, start=1):
            if col_idx == creation_col_idx:
                continue
            name = col_info["name"]
            if name in df_filtered.columns and _is_date_column(df_filtered[name]):
                date_col_idxs.add(col_idx)

        # Pré-conversion vectorisée des colonnes de date (perf + évite de parser
        # cellule par cellule). Aligné par position sur les lignes de df_filtered.
        date_values = {}
        for col_idx, col_info in enumerate(columns, start=1):
            if col_idx in date_col_idxs or col_idx == creation_col_idx:
                name = col_info["name"]
                if name in df_filtered.columns:
                    date_values[col_idx] = _parse_dates(df_filtered[name]).to_numpy()

        # Écrire les en-têtes de colonne avec style
        for col_idx, col_info in enumerate(columns, start=1):
            col_name = col_info["name"]
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = self.header_font_color
            cell.fill = self.header_fill
            cell.alignment = self.center_wrap
            cell.border = self.border

            # Définir la largeur de la colonne
            col_letter = ws.cell(row=1, column=col_idx).column_letter
            width = column_widths[col_name]
            # Colonne "Création" : largeur compacte (date / heure sur deux lignes)
            if col_idx == creation_col_idx:
                width = 12
            ws.column_dimensions[col_letter].width = width

        # En-tête de la colonne "Year" (ajoutée en fin de tableau)
        if year_col_idx:
            ycell = ws.cell(row=1, column=year_col_idx, value="Year")
            ycell.font = self.header_font_color
            ycell.fill = self.header_fill
            ycell.alignment = self.center_wrap
            ycell.border = self.border
            year_letter = ws.cell(row=1, column=year_col_idx).column_letter
            ws.column_dimensions[year_letter].width = 10

        # Hauteur de la ligne d'en-tête : permet d'afficher les libellés longs sur
        # plusieurs lignes (le wrap_text seul ne suffit pas, la hauteur reste sinon figée)
        ws.row_dimensions[1].height = 30

        # Écrire les données avec style (seulement les colonnes sélectionnées)
        for row_idx, row in enumerate(df_filtered.itertuples(index=False), start=2):
            pos = row_idx - 2  # position 0-based de la ligne (pour les dates pré-converties)
            for col_idx, value in enumerate(row, start=1):
                number_format = None
                # Colonne "Environnement et phase" → ne garder que le premier mot
                if col_idx == env_phase_col_idx:
                    value = _first_word(value)
                # Colonne "Création" → vraie date, affichée date/heure sur deux lignes
                elif col_idx == creation_col_idx:
                    dt = date_values[col_idx][pos]
                    if pd.isna(dt):
                        value = None if pd.isna(value) else value  # garde le texte si non-date
                    else:
                        value = pd.Timestamp(dt).to_pydatetime()
                        number_format = excel_date_format_creation
                # Autres colonnes de type date → vraie date typée + format d'affichage
                elif col_idx in date_col_idxs:
                    dt = date_values[col_idx][pos]
                    if pd.isna(dt):
                        value = None if pd.isna(value) else value  # garde le texte si non-date
                    else:
                        value = pd.Timestamp(dt).to_pydatetime()
                        number_format = excel_date_format
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                if number_format:
                    cell.number_format = number_format
                # Valeur "Critique" dans la colonne Priorité → texte rouge en gras
                if col_idx == priority_col_idx and _normalize(value) in CRITICAL_VALUES:
                    cell.font = self.critical_font
                else:
                    cell.font = self.cell_font
                # "Clé"/"Résumé" gardent l'alignement à gauche, toutes les autres sont centrées
                cell.alignment = self.left_align if col_idx in left_align_col_idxs else self.center_wrap
                cell.border = self.border

            # Valeur de la colonne "Year" pour cette ligne
            if year_col_idx:
                yval = years[row_idx - 2] if (row_idx - 2) < len(years) else None
                ycell = ws.cell(row=row_idx, column=year_col_idx, value=yval)
                ycell.font = self.cell_font
                ycell.alignment = self.center_align
                ycell.border = self.border

        # Activer les filtres sur la ligne d'en-tête (toute la plage de données)
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        # Figer la ligne d'en-tête pour garder les titres visibles au défilement
        ws.freeze_panes = "A2"

        # Enregistrer le fichier Excel généré
        wb.save(output_path)

        return {
            "output_path": str(output_path),
            "total_rows": len(df_filtered),
            "total_columns": len(columns) + (1 if year_col_idx else 0),
            "file_size": output_path.stat().st_size
        }

    def calculate_column_widths(self, df: pd.DataFrame, column_names: list[str]) -> dict[str, int]:
        """Largeur des colonnes basée sur les DONNÉES (et non sur l'intitulé).

        L'intitulé revient automatiquement à la ligne (wrap_text) lorsqu'il est plus
        long que la colonne : c'est ce qui produit les en-têtes sur deux lignes.
        On garantit toutefois une largeur ≥ au plus long mot de l'intitulé pour qu'aucun
        mot ne soit coupé, et on borne la largeur pour éviter les colonnes trop étroites
        ou trop larges."""
        widths = {}

        for col_name in column_names:
            # Plus long mot de l'intitulé : évite de couper un mot au milieu
            longest_header_word = max((len(w) for w in str(col_name).split()), default=0)

            # Longueur maximale des valeurs (échantillon des 1000 premières lignes)
            max_value_length = 0
            if col_name in df.columns:
                for value in df[col_name].head(min(1000, len(df))):
                    value_str = str(value) if pd.notna(value) else ""
                    max_value_length = max(max_value_length, len(value_str))

            # Largeur = max(données, plus long mot de l'en-tête) + petite marge
            optimal_width = max(max_value_length, longest_header_word) + 2

            # Bornée entre 12 et 28 : les intitulés longs passent ainsi sur deux lignes
            widths[col_name] = max(12, min(28, optimal_width))

        return widths
